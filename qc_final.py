#!/usr/bin/env python3
"""MATEL final quality-control gate. Exits non-zero if any check fails."""
import html
import json
import os
import re
import sys

from docx import Document

ROOT = os.path.dirname(os.path.abspath(__file__))
DOCX = os.path.join(ROOT, "MATEL_Threat_Intelligence_Report.docx")
PDF = os.path.join(ROOT, "MATEL_Threat_Intelligence_Report.pdf")
DOC = "MATEL_Threat_Intelligence_Report.docx"

fails = []


def check(name, ok, detail=""):
    print(("PASS  " if ok else "FAIL  ") + name + (("  " + detail) if detail else ""))
    if not ok:
        fails.append(name)


inv = json.load(open(os.path.join(ROOT, "template_invariants.json")))
d = Document(DOCX)
s = d.sections[0]

# 1. template preservation
for key, got, exp in [
        ("page width", s.page_width, inv["page"]["w"]),
        ("page height", s.page_height, inv["page"]["h"]),
        ("left margin", s.left_margin, inv["page"]["l"]),
        ("right margin", s.right_margin, inv["page"]["r"]),
        ("top margin", s.top_margin, inv["page"]["t"]),
        ("bottom margin", s.bottom_margin, inv["page"]["b"]),
        ("Normal font", d.styles["Normal"].font.name, inv["normal_font"]["name"]),
        ("Normal size", d.styles["Normal"].font.size, inv["normal_font"]["size"]),
        ("Heading 1 size", d.styles["Heading 1"].font.size, inv["h1_size"]),
        ("Title size", d.styles["Title"].font.size, inv["title_size"])]:
    check("template: " + key, str(got) == str(exp), str(got))

text = "\n".join(p.text for p in d.paragraphs)
for t in d.tables:
    for r in t.rows:
        for c in r.cells:
            text += "\n" + c.text
missing_headings = [h for h in inv["headings"] if h not in text]
check("template: all original headings retained", not missing_headings, str(missing_headings))

# 2. no unresolved placeholders
bad = [m.group(0) for m in re.finditer(r"\[[^\]]{2,80}\]", text)
       if not re.match(r"\[(?:MITRE|CISA|ABUSE)-\d\]|\[\d+\]", m.group(0))]
check("no unresolved placeholders", not bad, str(bad[:5]))

# 3. citation / reference consistency
refs = json.load(open(os.path.join(ROOT, "refs_final.json")))
surnames = {html.unescape(r["authors"][0]).strip().rstrip(".").split(" ")[-1].lower() for r in refs}
items = []
for group in re.findall(r"\(([^()]{2,140}?)\)", text):
    for part in group.split(";"):
        m = re.match(r"^(.*?),\s*((?:19|20)\d\d)$", part.strip())
        if m:
            items.append(m.group(1).strip())
check(">=25 academic references", len(refs) >= 25, "%d" % len(refs))
check(">=25 in-text citation occurrences", len(items) >= 25, "%d" % len(items))
uncited = [x for x in surnames if not any(x in k.lower() for k in items)]
check("every reference cited in text", not uncited, str(uncited))
orphan = sorted({k for k in items if not any(x in k.lower() for x in surnames)})
check("no citation without a reference", not orphan, str(orphan))

# 4. DOI / OA claims are not overstated
val = json.load(open(os.path.join(ROOT, "refs_final.json")))
check("all refs have resolved DOI", all(r["crossref_match"] for r in val))
check("all refs have legal open access", all(r["open_access"] for r in val))
bad_claim = [r["id"] for r in val if not str(r["scopus_indexed"]).startswith("Not independently")]
check("no unverified Scopus/SINTA claim", not bad_claim, str(bad_claim))

# 5. threat-intelligence artefacts
ioc_rows = 0
for f in ("ta542_emotet_ioc.csv", "ta505_dridex_ioc.csv"):
    p = os.path.join(ROOT, "ioc", f)
    with open(p) as fh:
        ioc_rows += max(0, sum(1 for _ in fh) - 1)
check("actor IOC rows collected", ioc_rows > 0, "%d rows" % ioc_rows)
check("FIN6 file intentionally absent", not os.path.exists(os.path.join(ROOT, "ioc", "fin6_ioc.csv")))
ttp = json.load(open(os.path.join(ROOT, "ttp_mapping.json")))
check("TTP mapping present", len(ttp["ttp"]) >= 50, "%d rows" % len(ttp["ttp"]))
check("every TTP row carries evidence",
      all(r.get("evidence") and r.get("technique_id") for r in ttp["ttp"]))
actors = {r["actor"] for r in ttp["ttp"]}
check("all three actors mapped", actors >= {"TA542", "TA505", "FIN6"}, str(sorted(actors)))

# 6. figures and captions
caps = re.findall(r"Gambar \d+\.", text)
tcaps = re.findall(r"Tabel \d+\.", text)
check("figure captions present", len(caps) >= 4, "%d" % len(caps))
check("table captions present", len(tcaps) >= 3, "%d" % len(tcaps))
missing_png = [p for p in ("ioc-pipeline", "propagation", "campaign-timeline",
                           "ttp-coverage", "actor-relationship", "workflow-methodology")
               if not os.path.exists(os.path.join(ROOT, "figures", p + ".png"))]
check("all report figures exist", not missing_png, str(missing_png))

# 7. diagram validation receipts
dg = [f for f in os.listdir(os.path.join(ROOT, "figures")) if f.endswith(".finalize-summary.json")]
passed = 0
for f in dg:
    r = json.load(open(os.path.join(ROOT, "figures", f)))
    if r.get("ok") and all(v == "pass" for v in (r.get("gates") or {}).values()):
        passed += 1
check("archify diagrams passing all gates", passed >= 8, "%d/%d" % (passed, len(dg)))

# 8. safety
payload = [f for f in os.listdir(ROOT)
           if re.search(r"\.(exe|dll|bin|scr|ps1|vbs|jar|iso|lnk|xlsm)$", f, re.I)]
check("no payload/executable in repository root", not payload, str(payload))

# 9. PDF
check("PDF generated", os.path.exists(PDF), "%d bytes" % (os.path.getsize(PDF) if os.path.exists(PDF) else 0))
if os.path.exists(PDF):
    raw = open(PDF, "rb").read()
    m = re.search(rb"/Count\s+(\d+)", raw)
    check("PDF has pages", bool(m) and int(m.group(1)) > 5, m.group(1).decode() if m else "?")

print()
if fails:
    print("QC FAILED: %d check(s)" % len(fails))
    for f in fails:
        print("  -", f)
    sys.exit(1)
print("QC PASSED: all checks green")
